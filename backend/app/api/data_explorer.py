"""Internal, local-only data explorer: what data we hold per fund and the original AMC files behind it.
Unauthenticated; mounted only when DATA_EXPLORER_ENABLED is true (see main.py)."""

import io
import zipfile
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.data.portfolio_sources import PORTFOLIO_SOURCES
from app.data.source_files import resolve_stored
from app.data.ter import latest_ter
from app.db import get_db
from app.models import (Document, NavPoint, Scheme, SchemeDetail, SchemeHolding, SchemeMetric, SchemePerformance,
                        SchemeTer, SourceFile)

router = APIRouter(prefix="/data", tags=["data-explorer"])

DATASETS = {
    "nav": "NAV history",
    "details": "Scheme master details",
    "ter": "Expense ratio (TER)",
    "metrics": "Computed metrics",
    "holdings": "Portfolio holdings",
    "source_file": "Source file for latest holdings",
    "performance_feed": "AMFI performance feed",
}
MAX_MEMBER_BYTES = 50 * 1024 * 1024
MAX_COLS = 30


def _universe(db: Session) -> list[Scheme]:
    return list(db.scalars(select(Scheme).where(Scheme.in_universe.is_(True))))


def _ym(d: date) -> tuple[int, int]:
    return d.year, d.month


def _have(db: Session, funds: list[Scheme]) -> dict[str, set[int]]:
    """Per dataset, the universe codes that have it."""
    codes = {f.amfi_code for f in funds}

    def codes_in(col) -> set[int]:
        return {c for c in db.scalars(select(col).distinct()) if c in codes}

    holdings = {c: m for c, m in db.execute(select(SchemeHolding.amfi_code, func.max(SchemeHolding.disclosure_month))
                                            .group_by(SchemeHolding.amfi_code)) if c in codes}
    files = {(a, _ym(m)) for a, m in db.execute(select(SourceFile.amc_name, SourceFile.disclosure_month))}
    return {
        "nav": codes_in(NavPoint.amfi_code) | {f.amfi_code for f in funds if f.latest_nav is not None},
        "details": codes_in(SchemeDetail.amfi_code),
        "ter": set(latest_ter(db, sorted(codes), "regular")),
        "metrics": codes_in(SchemeMetric.amfi_code),
        "holdings": set(holdings),
        "source_file": {f.amfi_code for f in funds if f.amfi_code in holdings and (f.amc_name, _ym(holdings[f.amfi_code])) in files},
        "performance_feed": codes_in(SchemePerformance.amfi_code),
    }


@router.get("/coverage")
def coverage(db: Session = Depends(get_db)):
    funds = _universe(db)
    total = len(funds)
    have = _have(db, funds)
    datasets = [{"key": k, "label": DATASETS[k], "have": len(have[k]), "total": total,
                 "pct": round(100 * len(have[k]) / total, 1) if total else 0.0, "note": ""} for k in DATASETS]
    headline = round(sum(d["pct"] for d in datasets) / len(datasets), 1) if datasets else 0.0
    return {"headline_pct": headline, "completed_pct": headline, "pending_pct": round(100 - headline, 1),
            "universe": total, "datasets": datasets,
            "reference_documents": {"loaded": db.scalar(select(func.count()).select_from(Document)), "note": "not per fund"}}


def _reason(amc: str, key: str) -> str:
    if key in ("holdings", "source_file"):
        for name, src in PORTFOLIO_SOURCES.items():
            if name.lower() in amc.lower():
                return src["how_to_fetch"]
    return "not loaded yet"


@router.get("/coverage/missing")
def missing(dataset: str, amc: str | None = None, q: str | None = None, limit: int = Query(50, ge=1, le=500),
            offset: int = Query(0, ge=0), db: Session = Depends(get_db)):
    if dataset not in DATASETS:
        raise HTTPException(404, "unknown dataset")
    funds = _universe(db)
    have = _have(db, funds)[dataset]
    gaps = sorted((f for f in funds if f.amfi_code not in have), key=lambda f: (f.amc_name, f.name))
    by_amc: dict[str, int] = {}
    for f in gaps:
        by_amc[f.amc_name] = by_amc.get(f.amc_name, 0) + 1
    shown = [f for f in gaps if (not amc or f.amc_name == amc) and (not q or q.lower() in f.name.lower())]
    return {"dataset": dataset, "total_missing": len(gaps),
            "by_amc": [{"amc": a, "missing": n, "reason": _reason(a, dataset)} for a, n in sorted(by_amc.items(), key=lambda x: -x[1])],
            "funds": [{"amfi_code": f.amfi_code, "name": f.name, "amc": f.amc_name, "category": f.category_key}
                      for f in shown[offset:offset + limit]]}


def _fund(db: Session, code: int) -> Scheme:
    f = db.get(Scheme, code)
    if not f:
        raise HTTPException(404, "fund not found")
    return f


def _files(db: Session, fund: Scheme, month: date | None = None) -> list[SourceFile]:
    rows = db.scalars(select(SourceFile).where(SourceFile.amc_name == fund.amc_name)
                      .order_by(SourceFile.disclosure_month.desc(), SourceFile.filename)).all()
    return [r for r in rows if month is None or _ym(r.disclosure_month) == _ym(month)]


def _latest_holdings_month(db: Session, code: int) -> date | None:
    return db.scalar(select(func.max(SchemeHolding.disclosure_month)).where(SchemeHolding.amfi_code == code))


def _kv(d: dict) -> list[list]:
    return [[k, v if isinstance(v, (int, float, str, type(None))) else str(v)] for k, v in d.items()]


def _file_table(rows: list[SourceFile]) -> dict:
    cols = ["id", "filename", "disclosure_month", "size_bytes", "via", "source_url"]
    return {"columns": cols, "rows": [[r.id, r.filename, r.disclosure_month, r.size_bytes, r.via, r.source_url] for r in rows],
            "total": len(rows), "as_of": rows[0].disclosure_month if rows else None,
            "source": rows[0].source_url if rows else None}


def _dataset(db: Session, fund: Scheme, key: str, limit: int, offset: int) -> dict | None:
    """The fund's rows for a dataset, or None when it has none."""
    code = fund.amfi_code
    if key == "nav":
        total = db.scalar(select(func.count()).select_from(NavPoint).where(NavPoint.amfi_code == code))
        if total:
            rows = db.execute(select(NavPoint.nav_date, NavPoint.nav).where(NavPoint.amfi_code == code)
                              .order_by(NavPoint.nav_date.desc()).limit(limit).offset(offset)).all()
            return {"columns": ["date", "nav"], "rows": [list(r) for r in rows], "total": total, "as_of": rows[0][0] if rows else None,
                    "source": "AMFI NAV history"}
        if fund.latest_nav is None:
            return None
        return {"columns": ["date", "nav"], "rows": [[fund.latest_nav_date, fund.latest_nav]], "total": 1,
                "as_of": fund.latest_nav_date, "source": "AMFI NAVAll"}
    if key == "details":
        d = db.get(SchemeDetail, code)
        if not d:
            return None
        base = {c: getattr(d, c) for c in ("amc_full_name", "scheme_base_name", "scheme_type", "scheme_category", "nav_name",
                                           "min_amount", "launch_date", "closure_date", "isins")}
        rows = [[k, v if isinstance(v, (int, float, str, type(None))) else str(v)] for k, v in {**base, **(d.raw or {})}.items()]
        return {"columns": ["field", "value"], "rows": rows[offset:offset + limit], "total": len(rows), "as_of": d.updated_at.date(),
                "source": "AMFI scheme master"}
    if key == "metrics":
        m = db.scalar(select(SchemeMetric).where(SchemeMetric.amfi_code == code).order_by(SchemeMetric.as_of.desc()).limit(1))
        if not m:
            return None
        rows = _kv({**(m.metrics or {}), "composite_score": m.composite_score, "tier": m.tier, "model_version": m.model_version})
        return {"columns": ["metric", "value"], "rows": rows[offset:offset + limit], "total": len(rows), "as_of": m.as_of,
                "source": "computed from NAV history"}
    if key == "ter":
        t = latest_ter(db, [code], "regular").get(code)
        if not t:
            return None
        cols = ["effective_from", "r_ber", "r_brokerage", "r_transaction", "r_statutory", "r_ter", "r_6a_b", "r_6a_c"]
        q = select(SchemeTer).where(SchemeTer.nsdl_code == t["nsdl_code"])
        total = db.scalar(select(func.count()).select_from(q.subquery()))
        rows = db.scalars(q.order_by(SchemeTer.effective_from.desc()).limit(limit).offset(offset)).all()
        return {"columns": cols, "rows": [[getattr(r, c) for c in cols] for r in rows], "total": total,
                "as_of": rows[0].effective_from if rows else None, "source": "AMFI TER disclosure (regular plan)"}
    if key == "holdings":
        month = _latest_holdings_month(db, code)
        if not month:
            return None
        q = select(SchemeHolding).where(SchemeHolding.amfi_code == code, SchemeHolding.disclosure_month == month)
        total = db.scalar(select(func.count()).select_from(q.subquery()))
        rows = db.scalars(q.order_by(SchemeHolding.weight_pct.desc()).limit(limit).offset(offset)).all()
        cols = ["security_name", "security_isin", "asset_class", "sector", "weight_pct", "market_value_inr", "quantity", "rating"]
        return {"columns": cols, "rows": [[getattr(r, c) for c in cols] for r in rows], "total": total, "as_of": month,
                "source": next((r.source_url for r in rows if r.source_url), "AMC monthly portfolio disclosure")}
    if key in ("source_file", "files"):
        month = _latest_holdings_month(db, code) if key == "source_file" else None
        if key == "source_file" and not month:
            return None
        rows = _files(db, fund, month)
        return _file_table(rows) if rows else None
    if key == "performance_feed":
        q = select(SchemePerformance).where(SchemePerformance.amfi_code == code)
        total = db.scalar(select(func.count()).select_from(q.subquery()))
        if not total:
            return None
        rows = db.scalars(q.order_by(SchemePerformance.as_of.desc()).limit(limit).offset(offset)).all()
        cols = ["as_of", "benchmark", "riskometer", "aum_cr", "returns", "benchmark_returns"]
        return {"columns": cols, "rows": [[r.as_of, r.benchmark, r.riskometer, r.aum_cr, str(r.returns), str(r.benchmark_returns)] for r in rows],
                "total": total, "as_of": rows[0].as_of, "source": "AMFI performance feed"}
    raise HTTPException(404, "unknown dataset")


@router.get("/funds/{amfi_code}/datasets")
def fund_datasets(amfi_code: int, db: Session = Depends(get_db)):
    fund = _fund(db, amfi_code)
    out = []
    for key in [*DATASETS, "files"]:
        d = _dataset(db, fund, key, 1, 0)
        out.append({"key": key, "label": DATASETS.get(key, "All source files for the AMC"), "available": d is not None,
                    "rows": d["total"] if d else 0, "as_of": d["as_of"] if d else None, "source": d["source"] if d else None})
    return out


@router.get("/funds/{amfi_code}/datasets/{key}")
def fund_dataset(amfi_code: int, key: str, limit: int = Query(200, ge=1, le=2000), offset: int = Query(0, ge=0),
                 db: Session = Depends(get_db)):
    d = _dataset(db, _fund(db, amfi_code), key, limit, offset)
    if d is None:
        raise HTTPException(404, "not available")
    return d


@router.get("/funds/{amfi_code}/files")
def fund_files(amfi_code: int, db: Session = Depends(get_db)):
    return [{"id": r.id, "filename": r.filename, "disclosure_month": r.disclosure_month, "size_bytes": r.size_bytes,
             "via": r.via, "source_url": r.source_url} for r in _files(db, _fund(db, amfi_code))]


def _registered(db: Session, file_id: str):
    row = db.get(SourceFile, file_id)
    path = resolve_stored(row.stored_path) if row else None
    if not path:
        raise HTTPException(404, "file not found")
    return row, path


def _sheet_preview(data: bytes, filename: str, sheet: str | None, limit: int) -> dict:
    def cell(v) -> str:
        return "" if v is None else str(v)

    if filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        names = book.sheet_names()
        name = sheet or names[0]
        if name not in names:
            raise HTTPException(404, "sheet not found")
        ws = book.sheet_by_name(name)
        rows = [[cell(c) for c in ws.row_values(i)[:MAX_COLS]] for i in range(min(ws.nrows, limit))]
        return {"sheets": names, "sheet": name, "rows": rows}
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        name = sheet or wb.sheetnames[0]
        if name not in wb.sheetnames:
            raise HTTPException(404, "sheet not found")
        rows = []
        for r in wb[name].iter_rows(max_row=limit, max_col=MAX_COLS, values_only=True):
            row = [cell(c) for c in r]
            while row and not row[-1]:
                row.pop()
            rows.append(row)
        return {"sheets": wb.sheetnames, "sheet": name, "rows": rows}
    finally:
        wb.close()


@router.get("/files/{file_id}/preview")
def file_preview(file_id: str, sheet: str | None = None, member: str | None = None,
                 limit: int = Query(100, ge=1, le=1000), db: Session = Depends(get_db)):
    row, path = _registered(db, file_id)
    if row.filename.lower().endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            infos = {i.filename: i for i in z.infolist() if not i.is_dir()}
            if member is None:
                return {"members": [{"name": n, "size": i.file_size} for n, i in infos.items()]}
            info = infos.get(member)
            if not info or not member.lower().endswith((".xls", ".xlsx")):
                raise HTTPException(404, "member not found")
            if info.file_size > MAX_MEMBER_BYTES:
                raise HTTPException(413, "member too large to preview")
            return _sheet_preview(z.read(info), member, sheet, limit)
    if not row.filename.lower().endswith((".xls", ".xlsx")):
        raise HTTPException(404, "not previewable")
    return _sheet_preview(path.read_bytes(), row.filename, sheet, limit)


@router.get("/files/{file_id}/download")
def file_download(file_id: str, db: Session = Depends(get_db)):
    row, path = _registered(db, file_id)
    return FileResponse(path, filename=row.filename)


TABLES = {"nav": "nav_history", "details": "scheme_detail", "ter": "scheme_ter", "metrics": "scheme_metric",
          "holdings": "scheme_portfolio", "source_file": "source_file", "performance_feed": "scheme_performance"}


@router.get("/catalog")
def catalog(q: str | None = None, kind: str | None = Query(None, pattern="^(dataset|file)$"), dataset: str | None = None,
            amc: str | None = None, available: bool | None = None, limit: int = Query(50, ge=1, le=500),
            offset: int = Query(0, ge=0), db: Session = Depends(get_db)):
    """Library catalogue: every (fund, dataset) book and every stored source file, gaps included."""
    items: list[dict] = []
    ql = q.lower() if q else None
    if kind in (None, "dataset"):
        funds = [f for f in _universe(db) if (not amc or f.amc_name == amc) and (not ql or ql in f.name.lower() or ql in f.amc_name.lower())]
        funds.sort(key=lambda f: (f.amc_name, f.name))
        have = _have(db, funds)
        months = {c: m for c, m in db.execute(select(SchemeHolding.amfi_code, func.max(SchemeHolding.disclosure_month))
                                              .group_by(SchemeHolding.amfi_code))}
        for f in funds:
            for key in DATASETS:
                if (dataset and key != dataset) or (available is not None and (f.amfi_code in have[key]) != available):
                    continue
                ok = f.amfi_code in have[key]
                m = months.get(f.amfi_code)
                where = f"Local DB: {TABLES[key]}" + (f" ({m:%b %Y})" if key == "holdings" and m else "")
                items.append({"id": f"{f.amfi_code}:{key}", "kind": "dataset", "title": f"{f.name}: {DATASETS[key]}",
                              "fund": {"name": f.name, "amfi_code": f.amfi_code}, "amc": f.amc_name, "dataset_key": key,
                              "available": ok, "location": where if ok else "Not in the library",
                              "details": {} if ok else {"note": _reason(f.amc_name, key)},
                              "open_url": f"/api/data/funds/{f.amfi_code}/datasets/{key}"})
    if kind in (None, "file") and not dataset and available is not False:
        for r in db.scalars(select(SourceFile).order_by(SourceFile.amc_name, SourceFile.disclosure_month.desc(), SourceFile.filename)):
            if (amc and r.amc_name != amc) or (ql and ql not in r.filename.lower() and ql not in r.amc_name.lower()):
                continue
            items.append({"id": r.id, "kind": "file", "title": r.filename, "fund": None, "amc": r.amc_name, "dataset_key": "files",
                          "available": True, "location": f"Stored file: source_files/{r.stored_path}. Original: {r.source_url}",
                          "details": {"as_of": r.disclosure_month, "size_bytes": r.size_bytes, "via": r.via, "sha256": r.sha256},
                          "open_url": f"/api/data/files/{r.id}/preview"})
    page = items[offset:offset + limit]
    for it in page:  # fill rows/as_of/source for the visible page only
        if it["kind"] == "dataset" and it["available"]:
            d = _dataset(db, db.get(Scheme, it["fund"]["amfi_code"]), it["dataset_key"], 1, 0)
            if d:
                it["details"] = {"rows": d["total"], "as_of": d["as_of"], "source": d["source"]}
    return {"total": len(items), "items": page}
